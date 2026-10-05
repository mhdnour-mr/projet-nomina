from flask import Flask, render_template, request, redirect, url_for, send_file, jsonify, session
import sqlite3, os, json, io, hashlib
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

app = Flask(__name__)
BASE = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.environ.get("DATA_DIR", BASE)
os.makedirs(DATA_DIR, exist_ok=True)
DB_PATH = os.path.join(DATA_DIR, "nomina.db")
app.secret_key = os.environ.get("SECRET_KEY") or os.urandom(24).hex()

def login_required(f):
    from functools import wraps
    @wraps(f)
    def decorated(*args, **kwargs):
        if not session.get("logged_in"):
            return redirect(url_for("login"))
        return f(*args, **kwargs)
    return decorated

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    conn = get_db()
    c = conn.cursor()
    c.execute("""CREATE TABLE IF NOT EXISTS users (
        id INTEGER PRIMARY KEY AUTOINCREMENT, username TEXT UNIQUE, password TEXT)""")
    try:
        c.execute("INSERT INTO users (username, password) VALUES (?, ?)",
                  ("admin", hashlib.sha256("admin123".encode()).hexdigest()))
    except sqlite3.IntegrityError:
        pass
    c.execute("""CREATE TABLE IF NOT EXISTS form1_applicant (
        id INTEGER PRIMARY KEY AUTOINCREMENT, full_name TEXT, mother_name TEXT, gender TEXT,
        birth_date TEXT, birth_country TEXT, residence TEXT, religion TEXT, marital_status TEXT,
        children_count TEXT, health_status TEXT, injury_type TEXT, employment_status TEXT,
        ready_to_work TEXT, travel_allowed TEXT, address TEXT, phone TEXT, whatsapp TEXT, email TEXT,
        education TEXT, field TEXT, specialty TEXT, institution TEXT, graduation_rate TEXT,
        graduation_year TEXT, work_name TEXT, work_title TEXT, work_start TEXT, work_end TEXT,
        work_reason TEXT, course_type TEXT, course_level TEXT, course_duration TEXT, course_date TEXT,
        course_org TEXT, course_doc TEXT, signature TEXT, form_date TEXT, created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP)""")
    c.execute("""CREATE TABLE IF NOT EXISTS form2_technical (
        id INTEGER PRIMARY KEY AUTOINCREMENT, applicant_id INTEGER, full_name TEXT, form_number TEXT,
        form_date TEXT, motivation TEXT, teamwork TEXT, communication TEXT, problem_solving TEXT,
        professionalism TEXT, confidence TEXT, organization TEXT, values_alignment TEXT,
        tech_scores TEXT, strengths TEXT, weaknesses TEXT, recommendation TEXT, created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP)""")
    c.execute("""CREATE TABLE IF NOT EXISTS form3_administrative (
        id INTEGER PRIMARY KEY AUTOINCREMENT, applicant_id INTEGER, full_name TEXT, specialty TEXT,
        form_number TEXT, form_date TEXT, degree_match TEXT, experience_required TEXT,
        certs_exist TEXT, courses_exist TEXT, employment_status TEXT, flexible_start TEXT,
        flexible_hours TEXT, willing_to_travel TEXT, expected_salary TEXT, within_budget TEXT,
        motivation TEXT, teamwork TEXT, communication TEXT, appearance TEXT, confidence TEXT,
        organization TEXT, values_alignment TEXT, strengths TEXT, weaknesses TEXT, recommendation TEXT,
        committee_rec TEXT, security_status TEXT, created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP)""")
    conn.commit()
    conn.close()

init_db()

@app.route("/login", methods=["GET", "POST"])
def login():
    db = get_db()
    if request.method == "POST":
        u = request.form.get("username", "")
        p = request.form.get("password", "")
        row = db.execute("SELECT * FROM users WHERE username=? AND password=?",
                         (u, hashlib.sha256(p.encode()).hexdigest())).fetchone()
        db.close()
        if row:
            session["logged_in"] = True
            return redirect(url_for("index"))
        return render_template("login.html", error="بيانات خاطئة")
    db.close()
    return render_template("login.html")

@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))

@app.route("/")
def index():
    db = get_db()
    f1 = db.execute("SELECT COUNT(*) FROM form1_applicant").fetchone()[0]
    f2 = db.execute("SELECT COUNT(*) FROM form2_technical").fetchone()[0]
    f3 = db.execute("SELECT COUNT(*) FROM form3_administrative").fetchone()[0]
    db.close()
    return render_template("index.html", f1=f1, f2=f2, f3=f3)

# ---------- FORM 1 ----------
@app.route("/form1", methods=["GET", "POST"])
def form1():
    db = get_db()
    if request.method == "POST":
        data = {k: request.form.get(k, "") for k in [
            "full_name","mother_name","gender","birth_date","birth_country","residence",
            "religion","marital_status","children_count","health_status","injury_type",
            "employment_status","ready_to_work","travel_allowed","address","phone","whatsapp",
            "email","education","field","specialty","institution","graduation_rate",
            "graduation_year","work_name","work_title","work_start","work_end","work_reason",
            "course_type","course_level","course_duration","course_date","course_org",
            "course_doc","signature","form_date"]}
        # Validation
        errors = []
        if not data["full_name"] or len(data["full_name"]) < 3:
            errors.append("الاسم الثلاثي مطلوب (3 أحرف على الأقل)")
        if data["phone"] and not data["phone"].replace(" ","").isdigit():
            errors.append("رقم الهاتف أرقام فقط")
        if data["email"] and "@" not in data["email"]:
            errors.append("البريد إلكتروني غير صحيح")
        if data["children_count"] and not data["children_count"].isdigit():
            errors.append("عدد الأولاد أرقام فقط")
        if errors:
            db.close()
            return render_template("form1.html", errors=errors, data=data)
        db.execute("INSERT INTO form1_applicant (" + ",".join(data.keys()) + ") VALUES (" + ",".join(["?"]*len(data)) + ")", tuple(data.values()))
        db.commit()
        applicant_id = db.execute("SELECT last_insert_rowid()").fetchone()[0]
        db.close()
        return redirect(url_for("complete"))
    db.close()
    return render_template("form1.html")

# ---------- FORM 2 (committee-only) ----------
@app.route("/form2/<int:applicant_id>", methods=["GET", "POST"])
@login_required
def form2(applicant_id):
    db = get_db()
    if request.method == "POST":
        tech_scores = json.dumps({f"q{i}": request.form.get(f"tq{i}", "") for i in range(1,7)})
        data = (applicant_id, request.form.get("full_name",""), request.form.get("form_number",""),
                request.form.get("form_date",""), request.form.get("motivation",""),
                request.form.get("teamwork",""), request.form.get("communication",""),
                request.form.get("problem_solving",""), request.form.get("professionalism",""),
                request.form.get("confidence",""), request.form.get("organization",""),
                request.form.get("values_alignment",""), tech_scores,
                request.form.get("strengths",""), request.form.get("weaknesses",""),
                request.form.get("recommendation",""))
        db.execute("INSERT INTO form2_technical (applicant_id,full_name,form_number,form_date,motivation,teamwork,communication,problem_solving,professionalism,confidence,organization,values_alignment,tech_scores,strengths,weaknesses,recommendation) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", data)
        db.commit()
        db.close()
        return redirect(url_for("form3", applicant_id=applicant_id))
    app1 = db.execute("SELECT full_name FROM form1_applicant WHERE id=?", (applicant_id,)).fetchone()
    db.close()
    return render_template("form2.html", applicant_id=applicant_id, app1=app1)

# ---------- FORM 3 (committee-only) ----------
@app.route("/form3/<int:applicant_id>", methods=["GET", "POST"])
@login_required
def form3(applicant_id):
    db = get_db()
    if request.method == "POST":
        data = (applicant_id, request.form.get("full_name",""), request.form.get("specialty",""),
                request.form.get("form_number",""), request.form.get("form_date",""),
                request.form.get("degree_match",""), request.form.get("experience_required",""),
                request.form.get("certs_exist",""), request.form.get("courses_exist",""),
                request.form.get("employment_status",""), request.form.get("flexible_start",""),
                request.form.get("flexible_hours",""), request.form.get("willing_to_travel",""),
                request.form.get("expected_salary",""), request.form.get("within_budget",""),
                request.form.get("motivation",""), request.form.get("teamwork",""),
                request.form.get("communication",""), request.form.get("appearance",""),
                request.form.get("confidence",""), request.form.get("organization",""),
                request.form.get("values_alignment",""), request.form.get("strengths",""),
                request.form.get("weaknesses",""), request.form.get("recommendation",""),
                request.form.get("committee_rec",""), request.form.get("security_status",""))
        db.execute("""INSERT INTO form3_administrative
            (applicant_id,full_name,specialty,form_number,form_date,degree_match,experience_required,
             certs_exist,courses_exist,employment_status,flexible_start,flexible_hours,willing_to_travel,
             expected_salary,within_budget,motivation,teamwork,communication,appearance,confidence,
             organization,values_alignment,strengths,weaknesses,recommendation,committee_rec,security_status)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", data)
        db.commit()
        db.close()
        return redirect(url_for("applicant_detail", id=applicant_id))
    app1 = db.execute("SELECT full_name FROM form1_applicant WHERE id=?", (applicant_id,)).fetchone()
    db.close()
    return render_template("form3.html", applicant_id=applicant_id, app1=app1)

@app.route("/complete")
def complete():
    return render_template("complete.html")

# ---------- ADMIN SETTINGS ----------
@app.route("/admin/settings", methods=["GET", "POST"])
@login_required
def settings():
    db = get_db()
    msg = ""
    if request.method == "POST":
        cur = request.form.get("current", "")
        new = request.form.get("new", "")
        conf = request.form.get("confirm", "")
        row = db.execute("SELECT password FROM users WHERE username=?", ("admin",)).fetchone()
        if not row or row["password"] != hashlib.sha256(cur.encode()).hexdigest():
            msg = "❌ كلمة المرور الحالية غير صحيحة"
        elif len(new) < 6:
            msg = "❌ كلمة المرور الجديدة يجب أن تكون 6 خانات على الأقل"
        elif new != conf:
            msg = "❌ كلمتا المرور غير متطابقتين"
        else:
            db.execute("UPDATE users SET password=? WHERE username=?",
                       (hashlib.sha256(new.encode()).hexdigest(), "admin"))
            db.commit()
            msg = "✅ تم تحديث كلمة المرور بنجاح"
    db.close()
    return render_template("settings.html", msg=msg)


# ---------- EDIT APPLICANT (committee) ----------
@app.route("/edit/<int:id>", methods=["GET", "POST"])
@login_required
def edit_applicant(id):
    keys = ["full_name","mother_name","gender","birth_date","birth_country","residence",
            "religion","marital_status","children_count","health_status","injury_type",
            "employment_status","ready_to_work","travel_allowed","address","phone","whatsapp","email",
            "education","field","specialty","institution","graduation_rate","graduation_year",
            "work_name","work_title","work_start","work_end","work_reason","course_type",
            "course_level","course_duration","course_date","course_org","course_doc",
            "signature","form_date"]
    db = get_db()
    if request.method == "POST":
        data = {k: request.form.get(k, "") for k in keys}
        errors = []
        if not data["full_name"] or len(data["full_name"]) < 3:
            errors.append("الاسم الثلاثي مطلوب (3 أحرف على الأقل)")
        if data["phone"] and not data["phone"].replace(" ", "").isdigit():
            errors.append("رقم الهاتف أرقام فقط")
        if data["email"] and "@" not in data["email"]:
            errors.append("البريد إلكتروني غير صحيح")
        if errors:
            db.close()
            return render_template("form1.html", data=data, edit_id=id, errors=errors)
        set_clause = ",".join([f"{k}=?" for k in keys])
        db.execute(f"UPDATE form1_applicant SET {set_clause} WHERE id=?", tuple(data.values()) + (id,))
        db.commit()
        db.close()
        return redirect(url_for("applicant_detail", id=id))
    row = db.execute("SELECT * FROM form1_applicant WHERE id=?", (id,)).fetchone()
    db.close()
    return render_template("form1.html", data=row, edit_id=id)


# ---------- DASHBOARD ----------
@app.route("/dashboard")
@login_required
def dashboard():
    db = get_db()
    q = request.args.get("q", "").strip()
    base = """SELECT f1.*, 
      CASE WHEN f2.id IS NOT NULL THEN 1 ELSE 0 END AS tech_filled,
      CASE WHEN f3.id IS NOT NULL THEN 1 ELSE 0 END AS admin_filled
      FROM form1_applicant f1
      LEFT JOIN form2_technical f2 ON f2.applicant_id=f1.id
      LEFT JOIN form3_administrative f3 ON f3.applicant_id=f1.id
      """
    if q:
        rows = db.execute(base + "WHERE f1.full_name LIKE ? ORDER BY f1.created_at DESC", (f"%{q}%",)).fetchall()
    else:
        rows = db.execute(base + "ORDER BY f1.created_at DESC").fetchall()
    db.close()
    return render_template("dashboard.html", rows=rows, title="الاستمارات المعبأة", q=q)

@app.route("/view/<int:id>")
@login_required
def view_form1(id):
    db = get_db()
    row = db.execute("SELECT * FROM form1_applicant WHERE id=?", (id,)).fetchone()
    db.close()
    return render_template("view_form1.html", r=row)

@app.route("/applicant/<int:id>")
@login_required
def applicant_detail(id):
    db = get_db()
    a1 = db.execute("SELECT * FROM form1_applicant WHERE id=?", (id,)).fetchone()
    a2 = db.execute("SELECT * FROM form2_technical WHERE applicant_id=?", (id,)).fetchone()
    a3 = db.execute("SELECT * FROM form3_administrative WHERE applicant_id=?", (id,)).fetchone()
    db.close()
    return render_template("applicant_detail.html", a1=a1, a2=a2, a3=a3)

@app.route("/delete/<int:id>")
@login_required
def delete(id):
    db = get_db()
    db.execute("DELETE FROM form1_applicant WHERE id=?", (id,))
    db.execute("DELETE FROM form2_technical WHERE applicant_id=?", (id,))
    db.execute("DELETE FROM form3_administrative WHERE applicant_id=?", (id,))
    db.commit()
    db.close()
    return redirect(url_for("dashboard"))

# ---------- EXPORT EXCEL ----------
@app.route("/export/applicant/<int:id>")
@login_required
def export_applicant(id):
    db = get_db()
    a1 = db.execute("SELECT * FROM form1_applicant WHERE id=?", (id,)).fetchone()
    a2 = db.execute("SELECT * FROM form2_technical WHERE applicant_id=?", (id,)).fetchone()
    a3 = db.execute("SELECT * FROM form3_administrative WHERE applicant_id=?", (id,)).fetchone()
    db.close()
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "المتقدم"
    ws1.append(["الحقل","القيمة"])
    for k in a1.keys():
        ws1.append([k, a1[k] or ""])
    if a2:
        ws2 = wb.create_sheet("التخصيص الفني")
        ws2.append(["الحقل","القيمة"])
        for k in a2.keys():
            ws2.append([k, a2[k] or ""])
    if a3:
        ws3 = wb.create_sheet("الإداري")
        ws3.append(["الحقل","القيمة"])
        for k in a3.keys():
            ws3.append([k, a3[k] or ""])
    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    nm = a1["full_name"] if a1 else "applicant"
    safe = "".join(c if c not in r'\/:*?"<>|' else "_" for c in nm).strip()[:50] or "applicant"
    return send_file(buf, download_name=f"{safe}_applicant.xlsx", as_attachment=True,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.route("/export")
@login_required
def export():
    db = get_db()
    rows = db.execute("SELECT * FROM form1_applicant ORDER BY created_at DESC").fetchall()
    db.close()
    wb = Workbook()
    ws = wb.active
    ws.title = "استمارات البيانات"
    headers = ["الاسم الثلاثي","الجنس","تاريخ الميلاد","البلد الأصلي","محل الإقامة","الديانة",
               "الحالة الاجتماعية","الوضع الصحي","الوضع الوظيفي","الجاهزية للعمل","قابلية السفر",
               "التعليم","التخصص","المؤسسة","السنة","البريد","الهاتف","تاريخ التعبئة"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9E1F2")
    for i, row in enumerate(rows, 2):
        vals = [row["full_name"], row["gender"], row["birth_date"], row["birth_country"],
                row["residence"], row["religion"], row["marital_status"], row["health_status"],
                row["employment_status"], row["ready_to_work"], row["travel_allowed"],
                row["education"], row["specialty"], row["institution"], row["graduation_year"],
                row["email"], row["phone"], row["form_date"]]
        for col, v in enumerate(vals, 1):
            ws.cell(row=i, column=col, value=v or "")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, download_name="nomina_export.xlsx", as_attachment=True, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=True)
